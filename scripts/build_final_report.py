from __future__ import annotations

import json
import os
from pathlib import Path

os.environ.setdefault("MPLBACKEND", "Agg")

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
ASSETS = ROOT / "artifacts" / "report_assets"
SUMMARY = json.loads((ROOT / "artifacts" / "evaluation_summary.json").read_text(encoding="utf-8"))
ROBUSTNESS = json.loads((ROOT / "artifacts" / "robustness" / "hsv_rf_robustness.json").read_text(encoding="utf-8"))
GREEN = "24543A"
PALE = "EEF5F0"
GOLD = "E1B63F"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd")) or OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    if shading.getparent() is None:
        properties.append(shading)


def set_cell_margins(cell, value: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar")) or OxmlElement("w:tcMar")
    if margins.getparent() is None:
        properties.append(margins)
    for side in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            margins.append(node)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, GREEN)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            if row_index % 2 == 0:
                set_cell_shading(cells[index], PALE)
    for row_index, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.5)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def paragraph(doc: Document, text: str, *, bold_start: str | None = None):
    item = doc.add_paragraph()
    item.style = doc.styles["Normal"]
    if bold_start and text.startswith(bold_start):
        item.add_run(bold_start).bold = True
        item.add_run(text[len(bold_start):])
    else:
        item.add_run(text)
    return item


def bullet(doc: Document, text: str):
    return doc.add_paragraph(text, style="List Bullet")


def caption(doc: Document, text: str):
    item = doc.add_paragraph(text)
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER
    item.paragraph_format.keep_with_next = False
    for run in item.runs:
        run.italic = True
        run.font.size = Pt(9)


def add_figure(doc: Document, path: Path, text: str, width: float = 6.2):
    item = doc.add_paragraph()
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER
    item.paragraph_format.keep_with_next = True
    item.add_run().add_picture(str(path), width=Inches(width))
    caption(doc, text)


def build_robustness_chart() -> Path:
    path = ASSETS / "robustness.png"
    rows = ROBUSTNESS["conditions"]
    labels = [row["condition"].replace("_", " ").title() for row in rows]
    values = [row["accuracy"] * 100 for row in rows]
    fig, axis = plt.subplots(figsize=(8.7, 4.5))
    colours = ["#2F7D4A" if value >= 90 else "#E1B63F" if value >= 75 else "#B84A4A" for value in values]
    bars = axis.barh(labels[::-1], values[::-1], color=colours[::-1])
    axis.set_xlim(0, 100)
    axis.set_xlabel("Accuracy on balanced 200-image subset (%)")
    axis.grid(axis="x", alpha=0.2)
    for bar, value in zip(bars, values[::-1], strict=True):
        axis.text(value + 1, bar.get_y() + bar.get_height() / 2, f"{value:.1f}", va="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=220, facecolor="white")
    plt.close(fig)
    return path


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    title_properties = styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("BMDS2133 Image Processing Assignment  |  ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    winner = SUMMARY["results"][0]
    test = winner["test_metrics"]
    accuracy_ci = test["accuracy_ci95"]

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Automated Banana Ripeness and Surface Quality Assessment")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("BMDS2133 Image Processing Assignment Report").bold = True
    doc.add_paragraph("\n")
    add_figure(doc, ROOT / "assets" / "banana-hero.png", "Banana image used as the application identity graphic.", 4.8)
    details = doc.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details.add_run("Programme: ______________________________\n").bold = True
    details.add_run("Group members and student IDs: ______________________________\n")
    details.add_run("Lecturer: ______________________________\n")
    details.add_run("Submission date: ______________________________")
    page_break(doc)

    doc.add_heading("Executive Summary", level=1)
    paragraph(doc, f"A reproducible computer-vision system was developed to classify banana images as unripe, ripe, overripe or rotten while estimating visible blemished surface area. Four classical pipelines were compared under a fixed train, validation and held-out test protocol. Model selection used validation macro-F1 only. The selected HSV histogram and Random Forest pipeline achieved {test['accuracy']:.2%} test accuracy and {test['macro_f1']:.2%} macro-F1 on 562 unseen images. Its bootstrap 95% accuracy interval was {accuracy_ci[0]:.2%} to {accuracy_ci[1]:.2%}.")
    paragraph(doc, "The prototype provides image and batch analysis, an evaluation dashboard, bounded video-frame processing, probability profiles, visual preprocessing evidence, CSV/PDF exports and input validation. Four-point homography supports perspective rectification. Statistical and robustness tests revealed stable performance under blur and small rotations but substantial sensitivity to strong sensor noise and illumination changes. Consequently, the classification results should be treated as decision support rather than a food-safety judgement. Segmentation and blemish accuracy remain conditional until the supplied annotation workflow is populated with genuine human labels.")

    doc.add_heading("Contents", level=1)
    for line in ["1 Introduction", "2 Literature Review", "3 Methodology", "4 Results and Discussion", "5 Prototype Implementation", "6 Conclusion", "References", "Appendix A Rubric Evidence Matrix", "Appendix B Group Contract", "Appendix C Generative AI Disclosure"]:
        paragraph(doc, line)
    page_break(doc)

    doc.add_heading("1 Introduction", level=1)
    doc.add_heading("1.1 Background and problem", level=2)
    paragraph(doc, "Banana colour and surface condition change during post-harvest ripening. Manual grading is rapid but subjective, particularly near adjacent maturity stages or when lighting changes the apparent peel colour. Non-destructive computer vision can standardise the observation process and retain visual evidence for review. Prior banana studies have shown that colour, texture and optical features can distinguish ripening stages, although accuracy depends on acquisition conditions and the definition of each class (Mazen & Nashat, 2019; Prabha & Kumar, 2015).")
    paragraph(doc, "The assignment problem was addressed as a four-class supervised classification task combined with interpretable surface analysis. Classical techniques were selected so that every feature and processing stage could be demonstrated and defended. The implementation also compares techniques, provides quantitative evaluation and integrates the best validation-selected method into a working dashboard.")
    doc.add_heading("1.2 SMART objectives", level=2)
    objectives = [
        "By the submission date, implement four distinct classical image-processing pipelines behind one tested interface for the four required ripeness classes.",
        "Using the fixed 11,793 training, 1,123 validation and 562 test images, select a model by validation macro-F1 and report final accuracy, precision, recall, F1-score, uncertainty and calibration metrics.",
        "By the final prototype review, provide single-image, batch and sampled-video analysis with visual preprocessing evidence and downloadable CSV/PDF results.",
        "Provide a reproducible workflow that measures segmentation IoU/Dice, bounding-box IoU and blemish error against at least 30 human-labelled images before quantitative surface claims are submitted.",
    ]
    for item in objectives:
        bullet(doc, item)
    doc.add_heading("1.3 Scope", level=2)
    paragraph(doc, "The scope covers visible-light JPG/PNG images and uploaded video files containing bananas. The system does not infer internal quality, chemical composition or food safety. Physical measurements require the user to supply a known fruit length; planar perspective correction requires four calibration-plane corners. These conditions are displayed in the interface so that uncalibrated pixel measurements are not presented as centimetres.")
    page_break(doc)

    doc.add_heading("2 Literature Review", level=1)
    doc.add_heading("2.1 Visual indicators of banana ripeness", level=2)
    paragraph(doc, "Banana grading studies consistently treat peel colour as a primary maturity cue. Mendoza and Aguilera (2004) used image analysis for ripening classification, while Prabha and Kumar (2015) combined colour and texture measurements for maturity assessment. Mazen and Nashat (2019) later reported an artificial neural network approach. These studies support the use of colour statistics, but they also imply a limitation: apparent colour changes with illumination, camera response and background reflection. A credible prototype therefore requires explicit robustness testing rather than a single headline accuracy.")
    doc.add_heading("2.2 Image enhancement and segmentation", level=2)
    paragraph(doc, "Median filtering suppresses impulse-like noise while preserving stronger edges. Contrast-limited adaptive histogram equalisation improves local contrast but restricts amplification in homogeneous regions (Pizer et al., 1987). In the present design, CLAHE is applied only to the L* channel so that chromatic ripeness cues are not deliberately shifted. The fruit mask combines border-colour distance in L*a*b* space with saturated HSV fruit colours, then applies morphological closing, opening and largest-component selection.")
    doc.add_heading("2.3 Feature and classifier choices", level=2)
    paragraph(doc, "HSV separates hue and saturation from brightness, making it suitable for interpretable green, yellow and brown ratios. L*a*b* represents perceptual colour differences and supports K-means clustering. Haralick et al. (1973) introduced grey-level co-occurrence statistics for texture description. Support-vector networks form high-dimensional decision boundaries (Cortes & Vapnik, 1995), while Random Forest combines decorrelated decision trees and can model non-linear feature interactions (Breiman, 2001). K-nearest neighbours provides a contrasting instance-based baseline.")
    add_table(doc, ["Approach", "Strength", "Limitation", "Role in this study"], [
        ["HSV ratios + SVM", "Interpretable colour and blemish features", "Thresholds react to lighting", "Explainable baseline"],
        ["L*a*b* K-means + KNN", "Explicit colour clusters", "K-means and neighbour cost", "Perceptual-colour comparator"],
        ["HSV histogram + Random Forest", "Compact global colour distribution", "Weak spatial information", "Validation-selected model"],
        ["Morphology + GLCM + Random Forest", "Shape and texture evidence", "Lower discrimination for colour-led stages", "Texture comparator"],
        ["Deep CNN", "Can learn spatial features", "Higher data, compute and explainability burden", "Reviewed but outside classical scope"],
    ], [1.35, 1.8, 1.8, 1.65])
    caption(doc, "Table 2.1 Literature-informed comparison of candidate approaches.")
    doc.add_heading("2.4 Evaluation design", level=2)
    paragraph(doc, "Accuracy alone can conceal poor performance in smaller classes. Macro-F1 gives equal importance to each class; confusion matrices show the direction of errors. Bootstrap confidence intervals quantify sampling uncertainty. Brier score, log loss and expected calibration error assess whether probabilities behave like meaningful confidence estimates. The test set must remain outside model selection; otherwise, the reported estimate becomes optimistically biased.")
    doc.add_heading("3 Methodology", level=1)
    add_figure(doc, ASSETS / "system_flow.png", "Figure 3.1 End-to-end system architecture and execution flow.")
    doc.add_heading("3.1 Dataset and integrity controls", level=2)
    paragraph(doc, "The Banana Ripeness Classification Dataset contains 13,478 images resized to 416 by 416 pixels and divided by its publisher into training, validation and test folders (Shahriar, 2025). The supplied split was retained. Filename-derived source identities and byte-identical SHA-256 content are checked across splits before training. Class counts are shown below.")
    counts = SUMMARY["split_counts"]
    add_table(doc, ["Split", "Overripe", "Ripe", "Rotten", "Unripe", "Total"], [[split.title(), *[str(counts[split][label]) for label in ("overripe", "ripe", "rotten", "unripe")], str(sum(counts[split].values()))] for split in ("train", "valid", "test")], [1.0, 1.05, 0.9, 1.0, 1.0, 1.0])
    caption(doc, "Table 3.1 Dataset distribution used without resampling the evaluation splits.")
    doc.add_heading("3.2 Preprocessing and surface measurements", level=2)
    for item in [
        "Decode the file, correct EXIF orientation, convert to 8-bit RGB and reject invalid dimensions.",
        "Standardise the geometry required by each feature extractor.",
        "Apply median denoising and luminance-only CLAHE for visual enhancement evidence.",
        "Estimate the foreground, refine it morphologically and retain the largest connected component.",
        "Extract its contour and bounding box, then divide brown/dark pixels by fruit-mask pixels to estimate the blemish ratio.",
    ]:
        bullet(doc, item)
    paragraph(doc, "Optional perspective correction orders four supplied plane corners, calculates a projective homography with OpenCV and warps the selected plane to a rectangle. Aspect-ratio-preserving object rectification remains available for visual comparison. If a known fruit length is supplied, pixels per centimetre are derived from the detected major dimension and are used to estimate projected area and contour perimeter.")
    doc.add_heading("3.3 Four comparison pipelines", level=2)
    add_table(doc, ["Pipeline", "Features", "Classifier", "Version"], [
        ["HSV segmentation", "Colour ratios, HSV statistics, area", "RBF SVM", "hsv-ratios-statistics-v1"],
        ["L*a*b* K-means", "Four centres, weights, mean, SD", "Distance-weighted KNN", "lab-kmeans4-v1"],
        ["HSV histogram", "16 bins per masked H, S and V", "200-tree Random Forest", "hsv-hist-16x3-masked-v2"],
        ["Morphology + GLCM", "Shape, intensity, four texture families", "250-tree Random Forest", "glcm-morphology-masked-v2"],
    ], [1.3, 2.25, 1.75, 1.55])
    caption(doc, "Table 3.2 Implemented algorithm definitions.")
    doc.add_heading("3.4 Training and selection protocol", level=2)
    paragraph(doc, "All estimators were fitted on training features. Candidate ranking used descending validation macro-F1, then validation accuracy. The winner was frozen and evaluated on the test split. Per-image class probabilities were retained for calibration analysis. A deterministic 1,000-resample bootstrap produced 95% intervals for accuracy and macro-F1.")
    doc.add_heading("3.5 Robustness and annotation protocols", level=2)
    paragraph(doc, "The validation-selected model was tested on the same balanced 200-image test subset under clean input, 30% dimming, 30% brightening, Gaussian blur, simulated sensor noise and a 10-degree rotation. These transformations were fixed before comparison. A separate annotation script calculates mask IoU, Dice, precision and recall, bounding-box IoU, plus blemish MAE, RMSE, bias and Pearson correlation. No segmentation metric is claimed until genuine labels are provided.")
    page_break(doc)

    doc.add_heading("4 Results and Discussion", level=1)
    doc.add_heading("4.1 Model comparison", level=2)
    add_table(doc, ["Rank", "Pipeline", "Validation macro-F1", "Test accuracy", "Test macro-F1"], [[str(index), item["algorithm"], f"{item['validation_metrics']['macro_f1']:.2%}", f"{item['test_metrics']['accuracy']:.2%}", f"{item['test_metrics']['macro_f1']:.2%}"] for index, item in enumerate(SUMMARY["results"], 1)], [0.6, 2.8, 1.35, 1.1, 1.1])
    caption(doc, "Table 4.1 Candidate ranking by validation performance with test results reported afterwards.")
    add_figure(doc, ASSETS / "test_metrics.png", "Figure 4.1 Held-out test accuracy and macro-F1 across the four pipelines.")
    paragraph(doc, f"HSV Histogram + Random Forest ranked first with {winner['validation_metrics']['macro_f1']:.2%} validation macro-F1. It achieved {test['accuracy']:.2%} accuracy and {test['macro_f1']:.2%} macro-F1 on the 562-image test set. The bootstrap accuracy interval ({accuracy_ci[0]:.2%} to {accuracy_ci[1]:.2%}) indicates the uncertainty expected from this test size. Its multiclass Brier score was {test['multiclass_brier_score']:.4f}, log loss was {test['log_loss']:.4f}, and expected calibration error was {test['expected_calibration_error']:.2%}.")
    doc.add_heading("4.2 Error analysis", level=2)
    add_figure(doc, ASSETS / "winner_confusion.png", "Figure 4.2 Confusion matrix for the validation-selected pipeline.", 5.1)
    report = test["classification_report"]
    add_table(doc, ["Class", "Precision", "Recall", "F1-score", "Support"], [[label.title(), f"{report[label]['precision']:.3f}", f"{report[label]['recall']:.3f}", f"{report[label]['f1-score']:.3f}", str(int(report[label]['support']))] for label in ("overripe", "ripe", "rotten", "unripe")], [1.4, 1.25, 1.25, 1.25, 1.0])
    caption(doc, "Table 4.2 Class-level held-out performance.")
    paragraph(doc, "The result supports the hypothesis that global colour distribution is more discriminative than the implemented texture-only representation for this dataset. The difference does not prove universal superiority. Histograms discard spot position and shape, and the dataset includes augmentation derived from a limited acquisition domain. The confidence error also shows that probability magnitudes should not be interpreted as guarantees.")
    doc.add_heading("4.3 Robustness experiment", level=2)
    robustness_chart = ASSETS / "robustness.png"
    add_figure(doc, robustness_chart, "Figure 4.3 Accuracy under fixed image-acquisition perturbations.")
    paragraph(doc, "Blur and a 10-degree rotation caused no accuracy loss on the selected subset. Dimming reduced accuracy by 21.0 percentage points, brightening reduced it by 14.0 points, and sensor noise reduced it by 39.5 points. These results identify illumination control and image-quality checking as priority improvements. They also prevent the clean test result from being overstated as general camera robustness.")
    doc.add_heading("4.4 Surface-measurement validity", level=2)
    paragraph(doc, "The prototype correctly computes a deterministic blemish-to-fruit pixel ratio and displays the underlying masks. However, algorithmic correctness is distinct from measurement validity. Human-labelled masks were not available in the repository, so IoU, Dice and blemish error cannot be reported honestly. The implemented annotation workflow makes this remaining evidence measurable. Excellent-category claims for segmentation accuracy should be made only after representative annotations are completed and the results are inserted.")
    page_break(doc)

    doc.add_heading("5 Prototype Implementation", level=1)
    add_figure(doc, ASSETS / "prototype_output.png", "Figure 5.1 Reproducible prototype output from a held-out image.")
    doc.add_heading("5.1 User interface and workflow", level=2)
    paragraph(doc, "The Streamlit interface separates assessment, model evaluation, model training and video analysis. Uploaded images are validated; duplicate filenames are disambiguated; every compatible model runs; and the validation-selected pipeline appears as the primary result. Confidence warnings, class-probability charts, segmentation overlays, bounding boxes and calibration evidence support manual review. Images, videos, charts and frames use bounded display widths to avoid page overflow.")
    doc.add_heading("5.2 Extra-effort functions", level=2)
    for item in [
        "Multiple-image processing with consolidated CSV and PDF exports.",
        "Uploaded-video decoding, bounded frame sampling, annotated frame playback, temporal confidence and surface-quality charts, and frame-level exports.",
        "Four-model training and evaluation dashboard with class-level confusion analysis.",
        "Four-point planar homography, optional physical scaling and aspect-ratio-preserving object views.",
        "Automated tests, model-schema validation, cached reproducible training, GitHub Actions and Docker deployment files.",
    ]:
        bullet(doc, item)
    doc.add_heading("5.3 Software quality", level=2)
    paragraph(doc, "Feature extractors implement a shared abstract pipeline contract. Preprocessing, evaluation, exports, model persistence, validation, robustness, video handling and charts are separated into modules. Versioned manifests reject stale models. Deterministic seeds and cached feature metadata support repeatability. The test suite covers feature shapes, preprocessing outputs, invalid files, geometry, metrics, leakage detection, PDF/CSV exports, video sampling and Streamlit start-up.")
    doc.add_heading("6 Conclusion", level=1)
    paragraph(doc, "The first objective was achieved: four classical pipelines were implemented and integrated. The second objective was achieved with validation-only selection and held-out evaluation; HSV Histogram + Random Forest produced 95.91% test macro-F1. The third objective was achieved through single-image, batch and sampled-video workflows with visual evidence and exports. The fourth objective is technically enabled but remains incomplete until the team supplies and reviews genuine annotations.")
    paragraph(doc, "The comparison demonstrates that masked global colour distributions were most effective on the supplied dataset, while GLCM and morphology alone were less discriminative. Robustness testing qualifies that conclusion because large illumination shifts and synthetic sensor noise caused substantial degradation. The most defensible next step is to complete annotation-based validation, then retrain with controlled illumination augmentation or add acquisition-quality rejection. External-camera and banana-variety testing should precede operational use.")

    doc.add_heading("References", level=1)
    references = [
        "Breiman, L. (2001). Random forests. Machine Learning, 45, 5-32. https://doi.org/10.1023/A:1010933404324",
        "Cortes, C., & Vapnik, V. (1995). Support-vector networks. Machine Learning, 20, 273-297. https://doi.org/10.1007/BF00994018",
        "Haralick, R. M., Shanmugam, K., & Dinstein, I. (1973). Textural features for image classification. IEEE Transactions on Systems, Man, and Cybernetics, SMC-3(6), 610-621. https://doi.org/10.1109/TSMC.1973.4309314",
        "Mazen, F. M. A., & Nashat, A. A. (2019). Ripeness classification of bananas using an artificial neural network. Arabian Journal for Science and Engineering, 44, 6901-6910. https://doi.org/10.1007/s13369-018-03695-5",
        "Mendoza, F., & Aguilera, J. M. (2004). Application of image analysis for classification of ripening bananas. Journal of Food Science, 69(9), E471-E477. https://doi.org/10.1111/j.1365-2621.2004.tb09932.x",
        "Pizer, S. M., Amburn, E. P., Austin, J. D., Cromartie, R., Geselowitz, A., Greer, T., ter Haar Romeny, B. M., Zimmerman, J. B., & Zuiderveld, K. J. (1987). Adaptive histogram equalization and its variations. Computer Vision, Graphics, and Image Processing, 39(3), 355-368. https://doi.org/10.1016/S0734-189X(87)80186-X",
        "Prabha, D. S., & Kumar, J. S. (2015). Assessment of banana fruit maturity by image processing technique. Journal of Food Science and Technology, 52, 1316-1327. https://doi.org/10.1007/s13197-013-1188-3",
        "Shahriar, S. M. (2025). Banana Ripeness Classification Dataset [Data set]. Kaggle. https://www.kaggle.com/datasets/shahriar26s/banana-ripeness-classification-dataset",
    ]
    for reference in references:
        item = paragraph(doc, reference)
        item.paragraph_format.left_indent = Inches(0.3)
        item.paragraph_format.first_line_indent = Inches(-0.3)
    page_break(doc)

    doc.add_heading("Appendix A Rubric Evidence Matrix", level=1)
    add_table(doc, ["Requirement", "Repository evidence", "Status"], [
        ["Noise removal and enhancement", "Median filter and L*-channel CLAHE", "Implemented"],
        ["Calibration and rectification", "Known-length scaling plus four-point homography", "Implemented; user input required"],
        ["Object detection", "Largest contour and bounding box", "Implemented; annotation validation pending"],
        ["Fruit classification", "Four pipelines, validation selection and held-out test", "Implemented"],
        ["Blemish relative to surface", "Brown/dark mask divided by fruit mask", "Implemented; ground-truth validation pending"],
        ["Dashboard and bulk GUI", "Assessment, evaluation, training, video and exports", "Implemented"],
        ["Video processing", "Sampled uploaded-video frame analysis", "Implemented extra effort"],
        ["PDF reporting", "Image and video PDF exports plus this report", "Implemented extra effort"],
        ["Team understanding", "Live explanation and code modification", "Must be demonstrated by students"],
    ], [1.75, 3.65, 1.55])
    caption(doc, "Table A.1 Evidence and conditions for final rubric claims.")
    page_break(doc)

    doc.add_heading("Appendix B Group Contract", level=1)
    paragraph(doc, "Complete the official institution form if one has been supplied. This working record should be replaced or signed before submission.")
    add_table(doc, ["Member and ID", "Agreed role and contribution", "Signature", "Date"], [["", "", "", ""] for _ in range(5)], [1.6, 3.25, 1.1, 1.0])
    paragraph(doc, "Agreement: Each member confirms that the submitted work has been reviewed, sources have been acknowledged, test results have not been fabricated, and every member can explain and modify the code allocated to them.")
    page_break(doc)

    disclosure_heading = doc.add_paragraph()
    disclosure_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    disclosure_heading.paragraph_format.left_indent = Inches(0)
    disclosure_heading.paragraph_format.first_line_indent = Inches(0)
    disclosure_heading.paragraph_format.space_after = Pt(8)
    disclosure_run = disclosure_heading.add_run("Appendix C Generative AI Disclosure")
    disclosure_run.bold = True
    disclosure_run.font.name = "Aptos Display"
    disclosure_run.font.size = Pt(18)
    paragraph(doc, "Generative AI tool: OpenAI Codex.")
    paragraph(doc, "Purpose: requirements review, debugging, implementation assistance, test design, documentation drafting and presentation preparation.")
    paragraph(doc, "Human verification required before submission: every generated or modified code path must be understood by the group; all numerical claims must be checked against repository artefacts; citations must be verified; member details and the institution's official disclosure fields must be completed.")
    paragraph(doc, "Representative prompt summary: The system was requested to be reviewed against the assignment specification, corrected where it failed, expanded with video processing and improved to meet the Excellent rubric. Subsequent prompts requested fixes for runtime exceptions and layout overflow.")

    output = OUT / "Banana_Ripeness_Assessment_Final_Report.docx"
    doc.save(output)
    print(output)
    return output


if __name__ == "__main__":
    build()
