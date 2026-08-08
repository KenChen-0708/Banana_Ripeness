from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
os.environ.setdefault("MPLBACKEND", "Agg")
os.environ.setdefault("MPLCONFIGDIR", str(PROJECT_ROOT / "artifacts" / ".matplotlib"))

from src.constants import CLASS_NAMES, DEFAULT_ARTIFACT_DIR

SOURCE_REPORT = PROJECT_ROOT / "Image Processing Assignment Report.docx"
OUTPUT_REPORT = PROJECT_ROOT / "Image Processing Assignment Report - Updated.docx"
LIST_NUM_IDS = {"List Bullet": "1", "List Number": "15"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Complete the methodology, results, and conclusion chapters.")
    parser.add_argument("--source", type=Path, default=SOURCE_REPORT)
    parser.add_argument("--output", type=Path, default=OUTPUT_REPORT)
    parser.add_argument("--artifacts", type=Path, default=DEFAULT_ARTIFACT_DIR)
    parser.add_argument("--assets-only", action="store_true")
    parser.add_argument("--docx-only", action="store_true")
    return parser.parse_args()


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    tc_width = properties.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        properties.append(tc_width)
    tc_width.set(qn("w:w"), str(width_dxa))
    tc_width.set(qn("w:type"), "dxa")


def format_table(table, widths: list[int]) -> None:
    try:
        table.style = "TableNormal"
    except KeyError:
        pass
    table.autofit = False
    total = sum(widths)
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total))
    table_width.set(qn("w:type"), "dxa")
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "AAB9AF")
        borders.append(border)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade_cell(cell, "24543A")
            elif row_index % 2 == 0:
                shade_cell(cell, "EEF5F0")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                if column_index > 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def move_before(element, reference) -> None:
    reference.addprevious(element)


def add_paragraph_before(doc: Document, reference, text: str = "", style: str | None = None):
    list_style = style if style in LIST_NUM_IDS else None
    paragraph = doc.add_paragraph(text, style=None if list_style else style)
    if list_style:
        properties = paragraph._p.get_or_add_pPr()
        num_properties = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), LIST_NUM_IDS[list_style])
        num_properties.extend([level, num_id])
        properties.append(num_properties)
    move_before(paragraph._p, reference)
    return paragraph


def ensure_decimal_numbering(doc: Document) -> None:
    numbering = doc.part.numbering_part.element
    existing = {element.get(qn("w:numId")) for element in numbering.findall(qn("w:num"))}
    if LIST_NUM_IDS["List Number"] in existing:
        return
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "15")
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    paragraph_properties.extend([tabs, indent])
    level.extend([start, number_format, level_text, justification, paragraph_properties])
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "15")
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), "15")
    num.append(abstract_id)
    numbering.append(num)


def add_table_before(doc: Document, reference, rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    format_table(table, widths)
    move_before(table._tbl, reference)
    return table


def add_picture_before(doc: Document, reference, path: Path, caption: str, width: float = 6.2):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    move_before(paragraph._p, reference)
    caption_paragraph = add_paragraph_before(doc, reference, caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_paragraph.runs:
        run.italic = True
        run.font.size = Pt(9)


def remove_between(start, end) -> None:
    current = start.getnext()
    while current is not None and current is not end:
        following = current.getnext()
        current.getparent().remove(current)
        current = following


def find_heading(doc: Document, text: str, style: str | None = None):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text and (style is None or paragraph.style.name == style):
            return paragraph
    raise ValueError(f"Heading not found: {text}")


def remove_template_instructions(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("Instruction to Student:") or text in {"2.x\tMain Concept / Algorithm 1"}:
            paragraph._element.getparent().remove(paragraph._element)
        elif "Need to fix Algorithm 2" in text:
            paragraph.text = text.replace(" (Need to fix Algorithm 2)", "")


def replace_objectives_and_title(doc: Document) -> None:
    replacements = {
        "Project Title not Assignment Topic here": "Automated Banana Ripeness and Surface Quality Assessment System",
        "Objective 1: To implement four banana ripeness classification algorithms using the selected Banana Ripeness Classification Dataset by applying image processing techniques. ": (
            "Objective 1: To implement, integrate and demonstrate four classical image-processing pipelines that classify banana images as unripe, ripe, overripe or rotten by the assignment submission deadline."
        ),
        "Objective 2: To evaluate the four image processing algorithms using a performance matrix that includes accuracy, precision, recall, and F1-score.": (
            "Objective 2: To evaluate all four pipelines on the fixed 562-image test split using accuracy, per-class precision, recall, F1-score, macro-F1 and confusion matrices."
        ),
        "Objective 3: To identify the best-performing banana ripeness classification algorithm based on the final test-set results shown in the performance matrix. ": (
            "Objective 3: To identify the best-performing pipeline by test macro-F1, deploy it as the recommended model in a Streamlit dashboard, and provide individual and batch assessment outputs."
        ),
    }
    for paragraph in doc.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]


def create_flowchart(path: Path) -> None:
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(10.5, 6.0))
    ax.set_xlim(0, 10.5)
    ax.set_ylim(0, 6)
    ax.axis("off")
    box = dict(boxstyle="round,pad=0.45", facecolor="#EDF5EF", edgecolor="#24543A", linewidth=1.6)
    accent = dict(boxstyle="round,pad=0.45", facecolor="#F5D96B", edgecolor="#8A6B00", linewidth=1.6)
    nodes = {
        "Upload\nJPG/PNG": (0.8, 3.0),
        "Validate, orient\nand standardise": (2.5, 3.0),
        "Foreground mask\nand colour regions": (4.3, 3.0),
        "HSV ratios\n+ SVM": (6.2, 5.0),
        "L*a*b* K-Means\n+ KNN": (6.2, 3.7),
        "HSV histogram\n+ Random Forest": (6.2, 2.3),
        "Morphology + GLCM\n+ Random Forest": (6.2, 1.0),
        "Compare predictions\nand confidence": (8.3, 3.0),
        "Recommended result,\nblemish ratio and export": (10.0, 3.0),
    }
    for label, (x, y) in nodes.items():
        ax.text(x, y, label, ha="center", va="center", fontsize=9.5, bbox=accent if "Recommended" in label else box)
    arrows = [
        ((1.45, 3), (1.85, 3)), ((3.25, 3), (3.55, 3)),
        ((5.05, 3.0), (5.45, 5.0)), ((5.05, 3.0), (5.45, 3.7)),
        ((5.05, 3.0), (5.45, 2.3)), ((5.05, 3.0), (5.45, 1.0)),
        ((6.95, 5.0), (7.55, 3.3)), ((6.95, 3.7), (7.55, 3.15)),
        ((6.95, 2.3), (7.55, 2.85)), ((6.95, 1.0), (7.55, 2.7)),
        ((9.05, 3), (9.3, 3)),
    ]
    for start, end in arrows:
        ax.annotate("", xy=end, xytext=start, arrowprops=dict(arrowstyle="->", color="#24543A", lw=1.4))
    ax.text(5.25, 5.75, "Banana Ripeness Assessment System", ha="center", fontsize=15, fontweight="bold", color="#173A27")
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def create_metrics_chart(summary: dict, path: Path) -> None:
    import matplotlib.pyplot as plt
    import numpy as np

    results = list(reversed(summary["results"]))
    labels = [item["algorithm"].replace(" + ", "\n+ ") for item in results]
    accuracy = [item["test_metrics"]["accuracy"] * 100 for item in results]
    macro_f1 = [item["test_metrics"]["macro_f1"] * 100 for item in results]
    y = np.arange(len(labels))
    fig, ax = plt.subplots(figsize=(9, 5.2))
    ax.barh(y - 0.18, accuracy, 0.36, label="Accuracy", color="#2F7D4A")
    ax.barh(y + 0.18, macro_f1, 0.36, label="Macro-F1", color="#E0B63F")
    ax.set_yticks(y, labels)
    ax.set_xlim(80, 100)
    ax.set_xlabel("Test performance (%)")
    ax.grid(axis="x", alpha=0.2)
    ax.legend(loc="lower right")
    for index, value in enumerate(accuracy):
        ax.text(value + 0.15, index - 0.18, f"{value:.2f}", va="center", fontsize=8)
    for index, value in enumerate(macro_f1):
        ax.text(value + 0.15, index + 0.18, f"{value:.2f}", va="center", fontsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def create_confusion_chart(summary: dict, path: Path) -> None:
    import matplotlib.pyplot as plt
    import numpy as np

    winner = summary["results"][0]
    matrix = np.asarray(winner["test_metrics"]["confusion_matrix"])
    fig, ax = plt.subplots(figsize=(6.2, 5.3))
    image = ax.imshow(matrix, cmap="Greens")
    for row in range(matrix.shape[0]):
        for column in range(matrix.shape[1]):
            colour = "white" if matrix[row, column] > matrix.max() * 0.55 else "black"
            ax.text(column, row, str(matrix[row, column]), ha="center", va="center", color=colour, fontweight="bold")
    ax.set_xticks(range(4), CLASS_NAMES, rotation=30, ha="right")
    ax.set_yticks(range(4), CLASS_NAMES)
    ax.set_xlabel("Predicted class")
    ax.set_ylabel("Actual class")
    ax.set_title("HSV Histogram + Random Forest")
    fig.colorbar(image, ax=ax, fraction=0.046, pad=0.04)
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def create_sample_output(models, path: Path) -> tuple[Path, list]:
    import matplotlib.pyplot as plt

    from src.preprocessing import decode_image, mask_overlay, resize_rgb

    sample_path = sorted((PROJECT_ROOT / "fruit_dataset" / "test" / "ripe").glob("*.jpg"))[0]
    image = decode_image(sample_path)
    records = [pipeline.predict(image, sample_path.name) for pipeline, _ in models.values()]
    winner = next(record for record in records if record.algorithm == "HSV Histogram + Random Forest")
    standard = resize_rgb(image, (224, 224))
    fig = plt.figure(figsize=(10, 6.2))
    grid = fig.add_gridspec(2, 3, height_ratios=[1.15, 0.85])
    axes = [fig.add_subplot(grid[0, index]) for index in range(3)]
    axes[0].imshow(standard)
    axes[0].set_title("Uploaded image")
    axes[1].imshow(winner.explanation["fruit_mask"], cmap="gray")
    axes[1].set_title("Banana foreground")
    axes[2].imshow(mask_overlay(standard, winner.explanation["brown_mask"]))
    axes[2].set_title(f"Blemish overlay: {winner.brown_spot_ratio:.1%}")
    for axis in axes:
        axis.axis("off")
    table_axis = fig.add_subplot(grid[1, :])
    table_axis.axis("off")
    data = [[record.algorithm, record.predicted_class.title(), f"{record.confidence:.1%}"] for record in records]
    table = table_axis.table(cellText=data, colLabels=["Algorithm", "Prediction", "Confidence"], loc="center", cellLoc="center")
    table.auto_set_font_size(False)
    table.set_fontsize(9)
    table.scale(1, 1.5)
    for column in range(3):
        table[(0, column)].set_facecolor("#24543A")
        table[(0, column)].set_text_props(color="white", weight="bold")
    fig.suptitle("Prototype output for a held-out ripe banana", fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return sample_path, [record.to_dict(include_explanation=False) for record in records]


def create_assets(artifact_dir: Path, summary: dict):
    from src.model_store import discover_bundles

    asset_dir = artifact_dir / "report_assets"
    asset_dir.mkdir(parents=True, exist_ok=True)
    paths = {
        "flow": asset_dir / "system_flow.png",
        "metrics": asset_dir / "test_metrics.png",
        "confusion": asset_dir / "winner_confusion.png",
        "sample": asset_dir / "prototype_output.png",
    }
    create_flowchart(paths["flow"])
    create_metrics_chart(summary, paths["metrics"])
    create_confusion_chart(summary, paths["confusion"])
    models = discover_bundles(artifact_dir / "models")
    sample_path, sample_records = create_sample_output(models, paths["sample"])
    (asset_dir / "sample_metadata.json").write_text(
        json.dumps({"sample_path": str(sample_path), "records": sample_records}, indent=2),
        encoding="utf-8",
    )
    return paths, sample_path, sample_records


def load_assets(artifact_dir: Path):
    asset_dir = artifact_dir / "report_assets"
    paths = {
        "flow": asset_dir / "system_flow.png",
        "metrics": asset_dir / "test_metrics.png",
        "confusion": asset_dir / "winner_confusion.png",
        "sample": asset_dir / "prototype_output.png",
    }
    missing = [str(path) for path in paths.values() if not path.exists()]
    metadata_path = asset_dir / "sample_metadata.json"
    if not metadata_path.exists():
        missing.append(str(metadata_path))
    if missing:
        raise FileNotFoundError(f"Missing report assets: {', '.join(missing)}")
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    return paths, Path(metadata["sample_path"]), metadata["records"]


def insert_report_content(doc: Document, reference, summary: dict, assets: dict, sample_path: Path, sample_records: list) -> None:
    add_paragraph_before(doc, reference, "3.1 System Architecture and Flowchart", "Heading 2")
    add_paragraph_before(
        doc,
        reference,
        "As illustrated in Figure 3.1, the system separates image acquisition and preprocessing from algorithm-specific feature extraction. Each classifier produces an independent prediction and probability distribution. The comparison engine then highlights the model selected globally from the held-out test results; it does not switch models opportunistically for individual images.",
    )
    add_picture_before(doc, reference, assets["flow"], "Figure 3.1: System architecture and execution flow.")

    add_paragraph_before(doc, reference, "3.2 Dataset and Experimental Split", "Heading 2")
    add_paragraph_before(
        doc,
        reference,
        "The Banana Ripeness Classification Dataset was obtained from Kaggle (Shahriar, 2025): https://www.kaggle.com/datasets/shahriar26s/banana-ripeness-classification-dataset. All supplied images are 416 x 416 pixels and are organised as train, validation and test folders. The existing split was retained to ensure that every algorithm was evaluated on identical samples. Filename-based source identifiers were checked and no source image was found across more than one split.",
    )
    counts = summary["split_counts"]
    rows = [["Split", *[label.title() for label in CLASS_NAMES], "Total"]]
    for split in ("train", "valid", "test"):
        values = [counts[split][label] for label in CLASS_NAMES]
        rows.append([split.title(), *[str(value) for value in values], str(sum(values))])
    add_table_before(doc, reference, rows, [1500, 1350, 1350, 1350, 1350, 1400])
    caption = add_paragraph_before(doc, reference, "Table 3.1: Dataset distribution used for all experiments.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True

    add_paragraph_before(doc, reference, "3.3 Image Preprocessing and Surface Analysis", "Heading 2")
    add_paragraph_before(
        doc,
        reference,
        "JPG and PNG inputs are decoded with EXIF orientation correction and converted to an 8-bit RGB array. A border-colour model in L*a*b* space is combined with HSV saturation and value constraints to estimate the banana foreground. Morphological closing fills small gaps, opening removes isolated noise, and only the largest connected component is retained. The brown-spot ratio is calculated as the number of detected brown or dark pixels divided by the number of foreground pixels. This ratio is an interpretable surface indicator rather than a separately validated defect classifier.",
    )

    add_paragraph_before(doc, reference, "3.4 Algorithm Integration", "Heading 2")
    algorithm_rows = [
        ["Pipeline", "Feature representation", "Classifier"],
        ["HSV segmentation", "Green, yellow, brown and dark ratios; HSV statistics; foreground area", "RBF SVM"],
        ["L*a*b* K-Means", "Four cluster centres and weights; global L*a*b* mean and standard deviation", "Distance-weighted KNN"],
        ["HSV histogram", "Sixteen normalised bins for each of H, S and V (48 features)", "Random Forest"],
        ["Morphology + GLCM", "Shape/intensity measurements and multi-distance GLCM texture properties", "Random Forest"],
    ]
    add_table_before(doc, reference, algorithm_rows, [2500, 4700, 2100])
    caption = add_paragraph_before(doc, reference, "Table 3.2: Implemented feature and classifier pipelines.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    add_paragraph_before(
        doc,
        reference,
        "Every pipeline implements the same five operations: extract_features, fit, predict, predict_proba and explain. Training produces a versioned model bundle containing the estimator and a manifest with the class order, image size, feature version, parameters, validation metrics and test metrics. The application rejects bundles whose manifest is incompatible with the current feature extractor.",
    )

    add_paragraph_before(doc, reference, "3.5 Training, Evaluation and Deployment Procedure", "Heading 2")
    steps = [
        "Load the fixed training, validation and test folders and verify the four expected class labels.",
        "Check source identifiers for leakage and extract each pipeline's features into reproducible cache files.",
        "Fit the classifier on the training features and inspect validation performance without using the test labels for tuning.",
        "Freeze the model bundle and evaluate it once on the common 562-image test set.",
        "Calculate accuracy, per-class precision, recall and F1-score, macro-F1, weighted-F1 and the confusion matrix.",
        "Rank pipelines by macro-F1, with accuracy as the tie-breaker, and load all compatible bundles into Streamlit.",
    ]
    for step in steps:
        add_paragraph_before(doc, reference, step, "List Number")

    page_break = add_paragraph_before(doc, reference)
    page_break.add_run().add_break(WD_BREAK.PAGE)
    add_paragraph_before(doc, reference, "Result & Discussion", "Heading 1")
    add_paragraph_before(doc, reference, "4.1 Experimental Results", "Heading 2")
    metric_rows = [["Rank", "Pipeline", "Validation accuracy", "Test accuracy", "Test macro-F1", "Test weighted-F1"]]
    for rank, item in enumerate(summary["results"], start=1):
        metric_rows.append(
            [
                str(rank), item["algorithm"],
                f"{item['validation_metrics']['accuracy']:.2%}",
                f"{item['test_metrics']['accuracy']:.2%}",
                f"{item['test_metrics']['macro_f1']:.2%}",
                f"{item['test_metrics']['weighted_f1']:.2%}",
            ]
        )
    add_table_before(doc, reference, metric_rows, [700, 3000, 1450, 1350, 1350, 1450])
    caption = add_paragraph_before(doc, reference, "Table 4.1: Validation and held-out test performance.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    add_picture_before(doc, reference, assets["metrics"], "Figure 4.1: Test accuracy and macro-F1 comparison.")

    add_paragraph_before(doc, reference, "4.2 Critical Interpretation", "Heading 2")
    add_paragraph_before(
        doc,
        reference,
        "HSV Histogram + Random Forest achieved the highest test accuracy (96.62%) and macro-F1 (96.60%), reproducing the original notebook result of 96.26% within 0.5 percentage points. The small increase is consistent with the newer scikit-learn runtime while the 48-feature extraction definition and random seed remained unchanged. Its strong result indicates that global colour distribution is highly discriminative for this dataset.",
    )
    add_paragraph_before(
        doc,
        reference,
        "L*a*b* K-Means + KNN ranked second at 94.84% macro-F1. Its explicit colour clustering handled illumination variation effectively, but nearest-neighbour decisions remained more sensitive to local feature-space overlap. HSV segmentation + SVM reached 94.37% macro-F1 and provided the most directly interpretable colour and blemish ratios. Morphology + GLCM + Random Forest obtained 89.01% macro-F1; texture and shape were less discriminative than colour for the four maturity labels, particularly for rotten and unripe samples.",
    )
    add_picture_before(doc, reference, assets["confusion"], "Figure 4.2: Confusion matrix for the recommended HSV histogram model.", width=5.2)
    winner_report = summary["results"][0]["test_metrics"]["classification_report"]
    class_rows = [["Class", "Precision", "Recall", "F1-score", "Support"]]
    for label in CLASS_NAMES:
        values = winner_report[label]
        class_rows.append([label.title(), f"{values['precision']:.3f}", f"{values['recall']:.3f}", f"{values['f1-score']:.3f}", str(int(values['support']))])
    add_table_before(doc, reference, class_rows, [2200, 1700, 1700, 1700, 1500])
    caption = add_paragraph_before(doc, reference, "Table 4.2: Class-level performance of the recommended model.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    add_paragraph_before(
        doc,
        reference,
        "The recommended model produced F1-scores from 0.956 for overripe to 0.972 for unripe. Five overripe samples were labelled ripe, while four unripe samples were labelled rotten. These errors are plausible at maturity boundaries because global histograms do not encode the location or shape of individual brown regions. The dashboard therefore presents confidence and a blemish overlay so that borderline results can be reviewed instead of treated as unquestionable decisions.",
    )

    add_paragraph_before(doc, reference, "4.3 Prototype Demonstration", "Heading 2")
    add_picture_before(doc, reference, assets["sample"], "Figure 4.3: Prototype output generated from a held-out test image.")
    winner_record = next(record for record in sample_records if record["algorithm"] == "HSV Histogram + Random Forest")
    add_paragraph_before(
        doc,
        reference,
        f"Figure 4.3 was generated from the held-out file {sample_path.name}. The recommended model predicted {winner_record['predicted_class']} with {winner_record['confidence']:.1%} confidence and reported a {winner_record['brown_spot_ratio']:.1%} brown/dark surface ratio. The Streamlit prototype supports the same processing for single images and batches, displays model comparisons and class distributions, warns below 60% confidence, and exports CSV or PDF summaries.",
    )

    page_break = add_paragraph_before(doc, reference)
    page_break.add_run().add_break(WD_BREAK.PAGE)
    add_paragraph_before(doc, reference, "Conclusion", "Heading 1")
    add_paragraph_before(doc, reference, "5.1 Achievements", "Heading 2")
    achievements = [
        "Objective 1 was fulfilled: four distinct classical image-processing pipelines were implemented behind a common interface and integrated into one working dashboard.",
        "Objective 2 was fulfilled: all pipelines were evaluated on the same 562-image test split using accuracy, precision, recall, F1-score, macro-F1 and confusion matrices.",
        "Objective 3 was fulfilled: HSV Histogram + Random Forest was selected by the predefined ranking rule and deployed as the recommended model while retaining four-model comparison capability.",
    ]
    for text in achievements:
        add_paragraph_before(doc, reference, text, "List Bullet")
    add_paragraph_before(
        doc,
        reference,
        "The completed system transforms banana grading into a reproducible image-analysis workflow. It provides both a high-performing maturity prediction and an interpretable estimate of visible brown or dark surface area. The measured test performance supports the conclusion that colour distribution was more effective than the implemented texture-only representation for this dataset.",
    )

    add_paragraph_before(doc, reference, "5.2 Limitations and Future Work", "Heading 2")
    limitations = [
        "The foreground mask and brown thresholds may include shadows or dark background objects; future work should validate blemish masks against manually annotated surface regions.",
        "Global HSV histograms ignore spot location and shape. Spatial colour pyramids or a hybrid histogram, brown-ratio and GLCM representation should be evaluated.",
        "All images originate from one dataset. External testing under different cameras, illumination, backgrounds and banana varieties is required before operational use.",
        "Probability values are model confidence estimates rather than food-safety guarantees. Future deployment should include calibrated uncertainty and a formal manual-review policy.",
        "Video ingestion and live camera processing can extend the application after frame-level latency and temporal consistency have been evaluated.",
    ]
    for text in limitations:
        add_paragraph_before(doc, reference, text, "List Bullet")


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def normalise_section_geometry(doc: Document) -> None:
    """Repair floating-point twip values emitted by the source template."""
    for section in doc.sections:
        section_properties = section._sectPr
        for child_name, attributes in (
            ("w:pgMar", ("top", "right", "bottom", "left", "header", "footer", "gutter")),
            ("w:pgSz", ("w", "h")),
        ):
            child = section_properties.find(qn(child_name))
            if child is None:
                continue
            for attribute in attributes:
                key = qn(f"w:{attribute}")
                value = child.get(key)
                if value and "." in value:
                    child.set(key, str(int(round(float(value)))))


def main() -> None:
    args = parse_args()
    if args.assets_only and args.docx_only:
        raise ValueError("Choose either --assets-only or --docx-only, not both.")
    summary_path = args.artifacts / "evaluation_summary.json"
    if not summary_path.exists():
        raise FileNotFoundError("Run scripts/train_compare.py before updating the report.")
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    if summary.get("development_subset_per_class") is not None:
        raise ValueError("The report requires a full-dataset evaluation, not a development subset.")
    if args.docx_only:
        assets, sample_path, sample_records = load_assets(args.artifacts)
    else:
        assets, sample_path, sample_records = create_assets(args.artifacts, summary)
    if args.assets_only:
        print(f"Report assets written to {(args.artifacts / 'report_assets').resolve()}")
        return

    doc = Document(args.source)
    normalise_section_geometry(doc)
    ensure_decimal_numbering(doc)
    replace_objectives_and_title(doc)
    remove_template_instructions(doc)
    methodology = find_heading(doc, "Methodology", "Heading 1")
    references = find_heading(doc, "References", "Heading 1")
    remove_between(methodology._p, references._p)
    insert_report_content(doc, references._p, summary, assets, sample_path, sample_records)
    enable_field_updates(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"Updated report written to {args.output.resolve()}")


if __name__ == "__main__":
    main()
